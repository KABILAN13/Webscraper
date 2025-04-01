import csv
import os
import sqlite3
from openpyxl import Workbook
from datetime import datetime

class DataStorage:
    @staticmethod
    def save_to_csv(data, filename):
        try:
        # Create 'outputs' directory if it doesn't exist
            os.makedirs('outputs', exist_ok=True)
        
        # Save to outputs folder by default
            filepath = os.path.join('outputs', filename)
        
            with open(filepath, 'w', newline='', encoding='utf-8') as f:
                writer = csv.DictWriter(f, fieldnames=data[0].keys())
                writer.writeheader()
                writer.writerows(data)
            return True
        except Exception as e:
            print(f"Error saving to CSV: {str(e)}")
            return False
    
    @staticmethod
    def save_to_excel(data, filename):
        if not data:
            return False
            
        try:
            wb = Workbook()
            ws = wb.active
            
            # Write headers
            headers = list(data[0].keys())
            ws.append(headers)
            
            # Write data
            for row in data:
                ws.append([row.get(header, '') for header in headers])
                
            wb.save(filename)
            return True
        except Exception as e:
            print(f"Error saving to Excel: {e}")
            return False
    
    @staticmethod
    def save_to_sqlite(data, db_file, table_name='products'):
        if not data:
            return False
            
        try:
            conn = sqlite3.connect(db_file)
            cursor = conn.cursor()
            
            # Create table if not exists
            sample = data[0]
            columns = ', '.join([f'"{k}" TEXT' for k in sample.keys()])
            cursor.execute(f'CREATE TABLE IF NOT EXISTS {table_name} ({columns}, scraped_at TIMESTAMP)')
            
            # Insert data
            for row in data:
                placeholders = ', '.join(['?'] * (len(row) + 1))
                columns = ', '.join([f'"{k}"' for k in row.keys()])
                values = list(row.values()) + [datetime.now()]
                cursor.execute(f'INSERT INTO {table_name} ({columns}, scraped_at) VALUES ({placeholders})', values)
            
            conn.commit()
            conn.close()
            return True
        except Exception as e:
            print(f"Error saving to SQLite: {e}")
            return False